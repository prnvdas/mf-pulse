"""Import a Groww "Mutual Funds holding statement" (.xlsx) from the local MF-Holding/ folder.

The statement is exact: it lists every folio with units, invested value, current value and XIRR. This script

  1. checks it against what the dashboard already holds (a wrong account or stale file is refused),
  2. syncs each fund's units and invested amount in docs/data/state.json to the statement's exact figures,
  3. writes config/folios.json: the per-folio breakdown WITHOUT any personal data. The statement itself
     contains your name, mobile number, PAN and full folio numbers; none of that is read into the output
     (folios are labelled by their last 4 digits only) and the MF-Holding/ folder is git-ignored.

Run it locally whenever you download a fresh statement (it cannot run in GitHub Actions, because the
statement must never be uploaded to a public repository):

    python src/import_statement.py [path/to/statement.xlsx] [--dry-run] [--force]
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import re
import sys

import openpyxl

from common import ROOT, load_portfolio, now_ist, read_json, write_json

STATEMENT_DIR = ROOT / "MF-Holding"


def norm(t: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(t).lower())


def num(x) -> float | None:
    try:
        return float(str(x).replace(",", "").replace("%", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return None


def find_header(ws) -> tuple[int, dict]:
    for r in range(1, ws.max_row + 1):
        row = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[r]]
        if "scheme name" in row and "units" in row:
            return r, {name: i for i, name in enumerate(row) if name}
    raise ValueError("couldn't find the holdings table (a row with 'Scheme Name' and 'Units')")


def read_statement(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Holdings"] if "Holdings" in wb.sheetnames else wb.worksheets[0]
    hr, col = find_header(ws)
    as_on = None
    summary = {}
    for r in range(1, hr):
        vals = [c.value for c in ws[r]]
        txt = str(vals[0] or "")
        m = re.search(r"HOLDINGS AS ON\s*(\d{4}-\d{2}-\d{2})", txt, re.I)
        if m:
            as_on = m.group(1)
        if str(vals[0] or "").strip().lower().startswith("total investments"):
            nxt = [c.value for c in ws[r + 1]]
            summary = {"invested": num(nxt[0]), "value": num(nxt[1]), "pl": num(nxt[2]), "xirr_pct": num(nxt[4])}
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        name = v[col["scheme name"]]
        if not name or num(v[col["units"]]) is None:
            continue
        rows.append({"scheme": str(name).strip(), "folio": str(v[col["folio no."]]).strip(), "units": num(v[col["units"]]),
                     "invested": num(v[col["invested value"]]), "value": num(v[col["current value"]]),
                     "xirr_pct": num(v[col["xirr"]]) if "xirr" in col else None,
                     "category": str(v[col["sub-category"]]) if "sub-category" in col and v[col["sub-category"]] else None})
    return {"as_on": as_on, "summary": summary, "rows": rows}


def match_fund(scheme: str, funds: list[dict]) -> dict | None:
    s = norm(scheme)
    best = None
    for f in funds:
        n = norm(f["name"])
        # the statement truncates long names; match on the common prefix
        k = min(len(s), len(n))
        if k >= 20 and (s[:k] == n[:k]):
            return f
        ov = len(set(re.findall(r"[a-z0-9]+", scheme.lower())) & set(re.findall(r"[a-z0-9]+", f["name"].lower())))
        if best is None or ov > best[0]:
            best = (ov, f)
    return best[1] if best and best[0] >= 4 else None


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", nargs="?", help="statement .xlsx (default: the newest in MF-Holding/)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true", help="skip the plausibility check against the current state")
    args = ap.parse_args()

    path = args.file or max(glob.glob(str(STATEMENT_DIR / "*.xlsx")), key=lambda p: __import__("os").path.getmtime(p), default=None)
    if not path:
        sys.exit("No statement found: put the Groww .xlsx in the MF-Holding/ folder.")
    st = read_statement(path)
    cfg = load_portfolio()
    state = read_json("state.json", {})
    print(f"[info] statement as on {st['as_on']}: {len(st['rows'])} folio rows (name, phone, PAN and full folio numbers are ignored)")

    per: dict[str, dict] = {}
    for r in st["rows"]:
        f = match_fund(r["scheme"], cfg["funds"])
        if not f:
            sys.exit(f"[error] '{r['scheme']}' doesn't match any fund in config/portfolio.yaml; nothing changed")
        p = per.setdefault(f["id"], {"fund_id": f["id"], "name": f["name"], "units": 0.0, "invested": 0.0, "value": 0.0, "folios": []})
        p["units"] += r["units"]; p["invested"] += r["invested"]; p["value"] += r["value"]
        p["folios"].append({"label": "…" + re.sub(r"\D", "", r["folio"])[-4:], "units": round(r["units"], 3), "invested": round(r["invested"], 2),
                            "value": round(r["value"], 2), "xirr_pct": r["xirr_pct"], "category": r["category"]})
    # ---- plausibility: refuse a statement that doesn't look like this portfolio
    problems = []
    for fid, p in per.items():
        s = state.get(fid, {})
        nav = s.get("last_nav")
        if s.get("units") and nav:
            have = s["units"] * nav
            if abs(p["value"] / have - 1) > 0.03:
                problems.append(f"{fid}: statement value ₹{p['value']:,.0f} vs dashboard ₹{have:,.0f}")
        inv = s.get("seed_invested")
        if inv and abs(p["invested"] / inv - 1) > 0.05:
            problems.append(f"{fid}: statement invested ₹{p['invested']:,.0f} vs dashboard ₹{inv:,.0f}")
    missing = [f["id"] for f in cfg["funds"] if f["id"] not in per]
    if missing:
        problems.append("funds missing from the statement: " + ", ".join(missing))
    if problems and not args.force:
        print("[refused] the statement doesn't match the dashboard (wrong account, or old file?). Use --force to override.\n  " + "\n  ".join(problems))
        sys.exit(1)

    print(f"\n{'fund':<24}{'units (stmt)':>14}{'units (dash)':>14}{'invested (stmt)':>17}{'invested (dash)':>17}")
    for fid, p in per.items():
        s = state.get(fid, {})
        print(f"{fid:<24}{p['units']:>14,.3f}{(s.get('units') or 0):>14,.3f}{p['invested']:>17,.2f}{(s.get('seed_invested') or 0):>17,.2f}")
        for fo in p["folios"]:
            print(f"   folio {fo['label']}: ₹{fo['invested']:,.0f} invested -> ₹{fo['value']:,.0f}, XIRR {fo['xirr_pct']}%")
    sm = st["summary"]
    print(f"\nstatement total: invested ₹{sm.get('invested'):,.2f}, value ₹{sm.get('value'):,.2f}, XIRR {sm.get('xirr_pct')}%")
    if args.dry_run:
        print("[dry-run] nothing written")
        return

    as_on = dt.date.fromisoformat(st["as_on"]) if st["as_on"] else now_ist().date()
    sip_day = {x["fund_id"]: int(x["day_of_month"]) for x in cfg.get("sips", [])}
    for fid, p in per.items():
        e = state.setdefault(fid, {})
        e["units"] = round(p["units"], 4)
        e["seed_invested"] = round(p["invested"], 2)
        # A statement dated before this month's SIP day does not contain that SIP yet: let reconcile add it later.
        if fid in sip_day and as_on.day < sip_day[fid] and e.get("sip_applied_month") == as_on.strftime("%Y-%m"):
            prev = (as_on.replace(day=1) - dt.timedelta(days=1)).strftime("%Y-%m")
            e["sip_applied_month"] = prev
            print(f"[info] {fid}: statement is dated before the SIP day ({sip_day[fid]}); this month's SIP will be added by reconcile")
    write_json("state.json", state)
    out = {"as_on": st["as_on"], "imported_at": now_ist().isoformat(), "source": "Groww holdings statement",
           "portfolio": {"invested": sm.get("invested"), "value": sm.get("value"), "xirr_pct": sm.get("xirr_pct")},
           "funds": [{"fund_id": p["fund_id"], "name": p["name"], "units": round(p["units"], 3), "invested": round(p["invested"], 2),
                      "value": round(p["value"], 2), "folios": p["folios"]} for p in per.values()]}
    import json
    (ROOT / "config" / "folios.json").write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n")
    print("[ok] state.json synced to the statement's units and invested amounts; config/folios.json written (no personal data)")


if __name__ == "__main__":
    main()
