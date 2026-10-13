"""Backend unit tests: the small pure pieces whose mistakes have cost real money-numbers before.

    python -m unittest discover -s tests -p "test_*.py"      (from the repo root; needs requirements.txt)
"""
import datetime as dt
import json
import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent / "src"))

import common  # noqa: E402
import detect_sips  # noqa: E402
import health  # noqa: E402
import import_transactions as itx  # noqa: E402
import estimate  # noqa: E402
import picks  # noqa: E402
import reconcile  # noqa: E402
import snapshot  # noqa: E402
import validate_data  # noqa: E402

D = dt.date


class ValidateData(unittest.TestCase):
    def test_nan_and_wrong_shapes_are_caught(self):
        with tempfile.TemporaryDirectory() as td:
            t = pathlib.Path(td)
            (t / "ok.json").write_text('{"a": 1}')
            (t / "nan.json").write_text('{"v": NaN}')
            (t / "inf.json").write_text('{"v": -Infinity}')
            (t / "history.json").write_text('{"not": "a list"}')
            (t / "portfolio_history.json").write_text('{"points": [{"d": "2026-01-01", "v": null, "i": 1}]}')
            (t / "broken.json").write_text("{")
            got = {p.name: validate_data.check(p) for p in sorted(t.glob("*.json"))}
        self.assertIsNone(got["ok.json"])
        for name in ("nan.json", "inf.json", "history.json", "portfolio_history.json", "broken.json"):
            self.assertIsNotNone(got[name], name)

    def test_every_committed_data_file_is_valid(self):
        bad = {p.name: w for p in validate_data.DATA.glob("*.json") if (w := validate_data.check(p))}
        self.assertEqual(bad, {})


class DetectSips(unittest.TestCase):
    def rows(self, fund, amounts_by_month):
        return [(D(y, m, day), fund, a) for (y, m, day), a in amounts_by_month]

    def test_amount_day_and_stamp_duty_rounding(self):
        rows = self.rows("f", [((2026, 7, 1), 49997.5), ((2026, 8, 3), 49997.5), ((2026, 9, 1), 49997.5), ((2026, 10, 1), 49997.5), ((2026, 10, 1), 499975.0)])
        sips, checked = detect_sips.detect(rows)
        self.assertEqual(sips, [{"fund_id": "f", "amount": 50000, "day_of_month": 1}])
        self.assertEqual(checked, ["f"])

    def test_old_amount_after_a_step_up_is_not_the_sip(self):
        rows = self.rows("f", [((2026, 4, 1), 24998.0), ((2026, 5, 4), 24998.0), ((2026, 6, 1), 24998.0), ((2026, 7, 1), 49997.0), ((2026, 8, 3), 49997.0), ((2026, 9, 1), 49997.0)])
        self.assertEqual(detect_sips.detect(rows)[0][0]["amount"], 50000)

    def test_stopped_sip_is_dropped_but_fund_is_still_checked(self):
        rows = self.rows("f", [((2026, 3, 1), 10000.0), ((2026, 4, 1), 10000.0), ((2026, 5, 1), 10000.0)]) + [(D(2026, 10, 5), "g", 1000.0)]
        sips, checked = detect_sips.detect(rows)
        self.assertEqual([s["fund_id"] for s in sips], [])
        self.assertEqual(checked, ["f", "g"])

    def test_one_off_lumpsums_are_not_sips(self):
        rows = self.rows("f", [((2026, 8, 5), 100000.0), ((2026, 9, 18), 500000.0), ((2026, 10, 1), 250000.0)])
        self.assertEqual(detect_sips.detect(rows)[0], [])


class Reconcile(unittest.TestCase):
    def test_sip_line_is_appended_once_and_only_if_the_file_exists(self):
        with tempfile.TemporaryDirectory() as td:
            old, reconcile.CONFIG = reconcile.CONFIG, pathlib.Path(td)
            try:
                reconcile.log_sip_transaction("x", D(2026, 11, 2), 50000, 300.1, 166.6)
                self.assertFalse((reconcile.CONFIG / "transactions.csv").exists())
                (reconcile.CONFIG / "transactions.csv").write_text("date,fund_id,amount,units,nav\n2026-10-01,x,1.00,1.0000,1")   # no trailing newline
                reconcile.log_sip_transaction("x", D(2026, 11, 2), 50000, 300.1, 166.6)
                reconcile.log_sip_transaction("x", D(2026, 11, 2), 50000, 300.1, 166.6)
                lines = (reconcile.CONFIG / "transactions.csv").read_text().splitlines()
            finally:
                reconcile.CONFIG = old
        self.assertEqual(lines[-1], "2026-11-02,x,50000.00,300.1000,166.6")
        self.assertEqual(len(lines), 3)


class Snapshot(unittest.TestCase):
    def test_non_finite_values_never_reach_the_file(self):
        with self.assertRaises(AssertionError):
            snapshot.point("2026-01-01", {"a": (float("nan"), 1.0)})
        p = snapshot.point("2026-01-01", {"a": (10.0, 8.0), "b": (5.0, 5.0)})
        self.assertEqual((p["v"], p["i"]), (15.0, 13.0))


class Health(unittest.TestCase):
    def w(self, f1, b1, r1, f3=None, b3=None, r3=None, of=20):
        return {"1Y": {"fund": f1, "bench": b1, "rank": r1, "of": of if r1 else None},
                "3Y": {"fund": f3, "bench": b3, "rank": r3, "of": of if r3 else None}, "5Y": {}}

    def test_verdicts(self):
        self.assertEqual(health.verdict(self.w(5, 3, 4, 14, 12, 3), 5)[0], "good")
        self.assertEqual(health.verdict(self.w(5, 3, 10, 14, 12, 12), 5)[1], "Beats its benchmark, mid-table")
        self.assertEqual(health.verdict(self.w(5, 3, 18, 14, 12, 19), 5)[1], "Beats its benchmark, trails its category")
        self.assertEqual(health.verdict(self.w(-2, 3, 18, 6, 12, 18), 5)[0], "weak")
        self.assertEqual(health.verdict(self.w(None, None, None), 0.5)[0], "new")

    def test_cagr_needs_enough_history(self):
        import pandas as pd
        s = pd.Series([100.0, 121.0], index=pd.to_datetime(["2024-01-01", "2026-01-01"]))
        self.assertAlmostEqual(health.cagr(s, pd.Timestamp("2026-01-01"), 2), 10.0, places=0)
        self.assertIsNone(health.cagr(s, pd.Timestamp("2026-01-01"), 5))


class ImportTransactions(unittest.TestCase):
    FUNDS = [{"id": "axis", "name": "Axis Small Cap Fund Direct Growth"}, {"id": "jio", "name": "JioBlackRock Flexi Cap Fund Direct Growth"}]

    def report(self, rows):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Name", "TEST USER"]); ws.append(["PAN", "AAAAA0000A"]); ws.append([])
        ws.append(["Scheme Name", "Transaction Type", "Units", "NAV", "Amount", "Date"])
        for r in rows:
            ws.append(r)
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); wb.save(f.name)
        return f.name

    def test_types_signs_and_strict_matching(self):
        path = self.report([
            ["Axis Small Cap Fund Direct Growth", "PURCHASE", 10, 100, 1000, dt.datetime(2024, 1, 5)],
            ["Axis Small Cap Fund Direct Growth", "REDEEM", 4, 125, 500, "10-Feb-2024"],
            ["Parag Parikh Flexi Cap Fund Direct Growth", "PURCHASE", 5, 60, 300, dt.datetime(2024, 2, 1)],   # must NOT land in JioBlackRock
            ["JioBlackRock Flexi Cap Fund Direct Growth", "PURCHASE", 3, 10, 30, dt.datetime(2025, 10, 13)],
            ["Axis Small Cap Fund Direct Growth", "PURCHASE", 1, 100, 100, dt.datetime(2024, 3, 1)],
        ])
        tx, skipped = itx.parse(path, {"funds": self.FUNDS})
        self.assertEqual([(t["fund_id"], t["amount"], t["units"]) for t in tx],
                         [("axis", 1000.0, 10.0), ("axis", -500.0, -4.0), ("axis", 100.0, 1.0), ("jio", 30.0, 3.0)])
        self.assertTrue(any(n.startswith("Parag Parikh") for n in skipped["unmatched"]))

    def test_a_statement_is_not_a_transaction_report(self):
        import openpyxl
        wb = openpyxl.Workbook(); wb.active.append(["Scheme Name", "Folio No.", "Units", "Invested Value"])
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); wb.save(f.name)
        self.assertFalse(itx.is_transaction_report(f.name))


class DataHealth(unittest.TestCase):
    NIFTY = [{"ticker": "^NSEI", "move_pct": 1.6}]
    FUND = [{"name": "A Fund Direct Growth", "coverage_pct": 90, "nav_move_pct": 1.2}]

    def test_clean_data_is_ok(self):
        h = estimate.data_health({"tickers": 100, "gap": 0, "repaired": 0, "unverified": 0, "outliers": [], "verifiable": True}, 100, 100, self.FUND, self.NIFTY)
        self.assertEqual((h["status"], h["issues"]), ("ok", []))

    def test_a_yahoo_gap_that_was_repaired_is_a_warning(self):
        h = estimate.data_health({"tickers": 264, "gap": 261, "repaired": 261, "unverified": 0, "outliers": [], "verifiable": True}, 264, 264, self.FUND, self.NIFTY)
        self.assertEqual(h["status"], "warn")
        self.assertIn("261", h["issues"][0]["text"])

    def test_unverifiable_closes_or_missing_prices_hold_the_estimate(self):
        self.assertEqual(estimate.data_health({"tickers": 100, "gap": 40, "repaired": 5, "unverified": 35, "outliers": [], "verifiable": True}, 100, 100, self.FUND, self.NIFTY)["status"], "bad")
        self.assertEqual(estimate.data_health({"tickers": 100, "verifiable": True}, 30, 100, self.FUND, self.NIFTY)["status"], "bad")

    def test_oddities_are_flagged(self):
        h = estimate.data_health({"tickers": 100, "gap": 0, "unverified": 0, "outliers": [("ABC.NS", -48.0)], "verifiable": True}, 99, 100,
                                 [{"name": "B Fund", "coverage_pct": 90, "nav_move_pct": -2.5}], self.NIFTY)
        texts = " ".join(i["text"] for i in h["issues"])
        self.assertEqual(h["status"], "warn")
        self.assertIn("ABC -48%", texts)                                        # split/bonus-style price, ignored
        self.assertIn("far from the Nifty", texts)                              # -2.5% against +1.6%
        self.assertEqual(estimate.data_health({"tickers": 1, "verifiable": True}, 1, 1, [], [])["issues"][0]["level"], "warn")   # no Nifty quote


class PreviousCloses(unittest.TestCase):
    def series(self, days, closes):
        import pandas as pd
        return pd.Series(closes, index=pd.to_datetime([f"2026-10-{d:02d}" for d in days]))

    def test_gap_in_yahoos_daily_bars_is_repaired_from_intraday(self):
        sessions = [D(2026, 10, 8), D(2026, 10, 9), D(2026, 10, 12), D(2026, 10, 13)]
        closes = {"GAP": self.series([8, 9, 13], [100.0, 101.0, 110.0]),          # the 12th is missing: 110 vs 101 would be a 2-day move
                  "FULL": self.series([8, 9, 12, 13], [50.0, 51.0, 52.0, 53.0]),
                  "STALE": self.series([8, 9, 12], [10.0, 11.0, 12.0])}           # no bar for the 13th yet
        old, estimate.intraday_close_on = estimate.intraday_close_on, lambda tk, day: {t: 105.0 for t in tk} if day == D(2026, 10, 12) else {}
        try:
            got = estimate.previous_closes(["GAP", "FULL", "STALE", "NONE"], closes, sessions)
        finally:
            estimate.intraday_close_on = old
        self.assertEqual(got["GAP"], 105.0)
        self.assertEqual(got["FULL"], 52.0)
        self.assertEqual(got["STALE"], 12.0)
        self.assertEqual(got["NONE"], 105.0)                                      # no daily data at all: also repaired

    def test_needs_two_sessions(self):
        self.assertEqual(estimate.previous_closes(["A"], {}, [D(2026, 10, 13)]), {})


class Picks(unittest.TestCase):
    def test_track_record_scores_old_picks_against_the_benchmark(self):
        today = dt.date.today()
        old = (today - dt.timedelta(days=35)).isoformat()
        young = (today - dt.timedelta(days=7)).isoformat()
        hist = [{"week_of": old, "generated_at": old, "crypto": [{"id": "a", "price": 10.0}, {"id": "b", "price": 10.0}], "stocks": [], "bench": {"btc": 100.0, "nifty500": 100.0}},
                {"week_of": young, "generated_at": young, "crypto": [{"id": "a", "price": 20.0}], "stocks": [], "bench": {"btc": 100.0, "nifty500": 100.0}}]
        rec = picks.track_record(hist, {"a": 15.0, "b": 8.0}, {}, 110.0, 120.0)
        first = [r for r in rec["rows"] if r["week_of"] == old][0]
        self.assertEqual((first["bench_pct"], first["n"]), (20.0, 2))
        self.assertEqual(first["avg_pct"], 15.0)        # (+50% and -20%) / 2
        self.assertEqual(first["beat"], 1)              # only 'a' beat Bitcoin's +20%
        self.assertEqual(rec["summary"]["crypto"]["weeks"], 1)    # the 7-day-old week is too young for the summary
        self.assertEqual(len([r for r in rec["rows"] if r["week_of"] == young]), 1)

    def test_helpers_survive_bad_numbers(self):
        self.assertIsNone(picks.fnum(float("nan")))
        self.assertIsNone(picks.fnum("x"))
        self.assertEqual(picks.fnum("3.5"), 3.5)
        import pandas as pd
        self.assertEqual(list(picks.pct_rank(pd.Series([1.0, 2.0, 3.0]))), [1 / 3, 2 / 3, 1.0])

    def test_the_crypto_name_filter_drops_copies_of_other_coins(self):
        for bad in ("wrapped-bitcoin Wrapped Bitcoin wbtc", "staked-ether Staked Ether steth", "usd-coin USDC usdc", "tether Tether usdt"):
            self.assertTrue(picks.EXCLUDE_NAME.search(bad), bad)
        for good in ("arbitrum Arbitrum arb", "near NEAR Protocol near"):
            self.assertFalse(picks.EXCLUDE_NAME.search(good), good)


class RealizedGains(unittest.TestCase):
    def test_fifo_split_between_long_and_short_term(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Scheme Name", "Transaction Type", "Units", "NAV", "Amount", "Date"])
        ws.append(["Some Equity Fund Direct Growth", "PURCHASE", 10, 100, 1000, dt.datetime(2024, 1, 10)])
        ws.append(["Some Equity Fund Direct Growth", "PURCHASE", 10, 120, 1200, dt.datetime(2024, 6, 10)])
        ws.append(["Some Equity Fund Direct Growth", "REDEEM", 15, 150, 2250, dt.datetime(2025, 2, 1)])
        ws.append(["Global NASDAQ FoF Direct Growth", "PURCHASE", 5, 10, 50, dt.datetime(2024, 1, 10)])
        ws.append(["Global NASDAQ FoF Direct Growth", "REDEEM", 5, 20, 100, dt.datetime(2025, 3, 1)])
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); wb.save(f.name)
        ev = itx.realized_events(f.name)
        eq = [e for e in ev if e["equity"]]
        self.assertEqual([(e["term"], e["gain"], e["fy"]) for e in eq], [("long", 500.0, "2024-25"), ("short", 150.0, "2024-25")])
        self.assertEqual([e["equity"] for e in ev if "NASDAQ" in e["scheme"]], [False])
        self.assertEqual(itx.fy_of(dt.date(2026, 3, 31)), "2025-26")
        self.assertEqual(itx.fy_of(dt.date(2026, 4, 1)), "2026-27")


class SoldFunds(unittest.TestCase):
    def test_a_closed_fund_and_its_sip_leave_the_portfolio_unless_asked_for(self):
        with tempfile.TemporaryDirectory() as td:
            t = pathlib.Path(td)
            (t / "portfolio.yaml").write_text("funds:\n  - {id: a, name: A}\n  - {id: b, name: B}\nsips:\n  - {fund_id: a, amount: 1, day_of_month: 1}\n  - {fund_id: b, amount: 2, day_of_month: 1}\n")
            (t / "funds_closed.json").write_text(json.dumps({"closed": [{"id": "b"}]}))
            old, common.CONFIG = common.CONFIG, t
            try:
                now, every = common.load_portfolio(), common.load_portfolio(include_closed=True)
            finally:
                common.CONFIG = old
        self.assertEqual([f["id"] for f in now["funds"]], ["a"])
        self.assertEqual([s["fund_id"] for s in now["sips"]], ["a"])
        self.assertEqual([f["id"] for f in every["funds"]], ["a", "b"])


class LoadPortfolioOverlay(unittest.TestCase):
    def test_learned_values_win_and_yaml_is_the_fallback(self):
        with tempfile.TemporaryDirectory() as td:
            t = pathlib.Path(td)
            (t / "portfolio.yaml").write_text("funds:\n  - {id: a, name: A, ter_annual_pct: 0.5}\n  - {id: b, name: B, ter_annual_pct: 0.4}\nsips:\n  - {fund_id: a, amount: 1000, day_of_month: 3}\n  - {fund_id: b, amount: 2000, day_of_month: 3}\n")
            (t / "fund_meta.json").write_text(json.dumps({"a": {"expense_ratio": 0.9}}))
            (t / "sips_detected.json").write_text(json.dumps({"funds_checked": ["a"], "sips": [{"fund_id": "a", "amount": 5000, "day_of_month": 1}]}))
            (t / "funds_auto.json").write_text(json.dumps([{"id": "c", "name": "C"}, {"id": "a", "name": "dup"}]))
            old, common.CONFIG = common.CONFIG, t
            try:
                cfg = common.load_portfolio()
            finally:
                common.CONFIG = old
        self.assertEqual([f["id"] for f in cfg["funds"]], ["a", "b", "c"])
        self.assertEqual({f["id"]: f["ter_annual_pct"] for f in cfg["funds"] if "ter_annual_pct" in f}, {"a": 0.9, "b": 0.4})
        self.assertEqual(sorted((s["fund_id"], s["amount"], s["day_of_month"]) for s in cfg["sips"]), [("a", 5000, 1), ("b", 2000, 3)])


if __name__ == "__main__":
    unittest.main()
